"""Short showcase PPT — WHITE THEME — 8 slides, black & blue text, mocked UI."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---------- palette: WHITE / BLACK / BLUE ----------
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
BLACK       = RGBColor(0x10, 0x14, 0x1A)          # near-black for headlines
TEXT_DARK   = RGBColor(0x1F, 0x29, 0x37)          # body text — soft black
TEXT_MUTED  = RGBColor(0x60, 0x6B, 0x7B)          # captions / sub-text
BLUE        = RGBColor(0x1D, 0x4E, 0xD8)          # primary blue
BLUE_DEEP   = RGBColor(0x14, 0x2A, 0x6B)          # deep blue for accents
BLUE_LIGHT  = RGBColor(0xE8, 0xEF, 0xFC)          # very light blue panel
BLUE_LINE   = RGBColor(0xC9, 0xD7, 0xF2)          # light blue border
BORDER      = RGBColor(0xE4, 0xE7, 0xEB)          # neutral border
PANEL_BG    = RGBColor(0xF7, 0xF9, 0xFC)          # off-white panel

# severity colors retained for meaning in mock UI only
SEV_CRIT  = RGBColor(0xDC, 0x26, 0x26)
SEV_HIGH  = RGBColor(0xEA, 0x58, 0x0C)
SEV_MED   = RGBColor(0xEA, 0xB3, 0x08)
SEV_LOW   = RGBColor(0x3B, 0x82, 0xF6)
SEV_OK    = RGBColor(0x10, 0xB9, 0x81)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height

# ---------- helpers ----------
def add_blank():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    # explicit white background (in case template default differs)
    bg = s.background.fill
    bg.solid(); bg.fore_color.rgb = WHITE
    return s

def rect(slide, x, y, w, h, fill, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb = line
    shp.shadow.inherit = False
    return shp

def rounded(slide, x, y, w, h, fill, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb = line
    shp.shadow.inherit = False
    try: shp.adjustments[0] = 0.1
    except Exception: pass
    return shp

def oval(slide, x, y, w, h, fill, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb = line
    shp.shadow.inherit = False
    return shp

def text(slide, x, y, w, h, body, *, size=14, bold=False, color=TEXT_DARK,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Calibri"):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    lines = body.split("\n") if isinstance(body, str) else body
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = font
    return tb

def page_header(s, title, subtitle=None):
    # thin blue accent strip at very top
    rect(s, 0, 0, SW, Inches(0.08), BLUE)
    # left blue stripe
    rect(s, 0, Inches(0.08), Inches(0.08), SH - Inches(0.08), BLUE)
    text(s, Inches(0.4), Inches(0.25), Inches(12.5), Inches(0.55),
         title, size=24, bold=True, color=BLACK)
    if subtitle:
        text(s, Inches(0.4), Inches(0.75), Inches(12.5), Inches(0.35),
             subtitle, size=12, color=BLUE)
    # underline
    rect(s, Inches(0.4), Inches(1.05), Inches(1.0), Inches(0.04), BLUE)
    # footer
    text(s, Inches(0.4), SH - Inches(0.3), Inches(12), Inches(0.22),
         "Multi-Cloud CSPM Platform  |  AWS • Azure • GCP",
         size=8, color=TEXT_MUTED)

# ---------- mock UI builders ----------
def browser_chrome(slide, x, y, w, h, url):
    rounded(slide, x, y, w, h, WHITE, line=BORDER)
    rect(slide, x + Inches(0.04), y + Inches(0.04), w - Inches(0.08), Inches(0.32), PANEL_BG)
    for i, c in enumerate([RGBColor(0xEF,0x46,0x46),
                           RGBColor(0xF5,0x9E,0x0B),
                           RGBColor(0x10,0xB9,0x81)]):
        oval(slide, x + Inches(0.12) + Inches(0.18) * i, y + Inches(0.12),
             Inches(0.12), Inches(0.12), c)
    rounded(slide, x + Inches(0.85), y + Inches(0.08), w - Inches(1.0), Inches(0.24),
            WHITE, line=BORDER)
    text(slide, x + Inches(0.95), y + Inches(0.1), w - Inches(1.2), Inches(0.2),
         url, size=8, color=TEXT_MUTED)

def severity_chip(slide, x, y, w, h, label, color):
    rounded(slide, x, y, w, h, color)
    text(slide, x, y + Inches(0.03), w, h - Inches(0.03),
         label, size=7, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

def mock_dashboard(slide, x, y, w, h):
    browser_chrome(slide, x, y, w, h, "cspm.local/dashboard")
    cy = y + Inches(0.45)
    text(slide, x + Inches(0.18), cy, Inches(4), Inches(0.3),
         "Security Posture", size=11, bold=True, color=BLACK)
    text(slide, x + Inches(0.18), cy + Inches(0.25), Inches(4), Inches(0.2),
         "Across 3 clouds, 12 accounts", size=7, color=TEXT_MUTED)
    kpis = [("87", "Posture",   BLUE),
            ("142","Findings",  BLUE_DEEP),
            ("9",  "Critical",  SEV_CRIT),
            ("31", "Resources", BLUE)]
    kx = x + Inches(0.18); ky = cy + Inches(0.6)
    kw = (w - Inches(0.45)) / 4 - Inches(0.05)
    for i, (n, l, c) in enumerate(kpis):
        rounded(slide, kx + (kw + Inches(0.05)) * i, ky, kw, Inches(0.7),
                WHITE, line=BORDER)
        rect(slide, kx + (kw + Inches(0.05)) * i, ky, Inches(0.05), Inches(0.7), c)
        text(slide, kx + (kw + Inches(0.05)) * i + Inches(0.12), ky + Inches(0.05),
             kw - Inches(0.12), Inches(0.32), n, size=15, bold=True, color=BLACK)
        text(slide, kx + (kw + Inches(0.05)) * i + Inches(0.12), ky + Inches(0.42),
             kw - Inches(0.12), Inches(0.25), l, size=7, color=TEXT_MUTED)
    cy2 = ky + Inches(0.85)
    rounded(slide, x + Inches(0.18), cy2, w - Inches(0.36), h - (cy2 - y) - Inches(0.18),
            PANEL_BG, line=BORDER)
    text(slide, x + Inches(0.3), cy2 + Inches(0.08), Inches(3), Inches(0.2),
         "Findings trend (30d)", size=8, bold=True, color=BLACK)
    base_y = cy2 + h - (cy2 - y) - Inches(0.3)
    bx = x + Inches(0.3)
    bw = (w - Inches(0.6)) / 14
    heights = [0.25,0.35,0.3,0.5,0.4,0.45,0.6,0.5,0.7,0.55,0.45,0.6,0.5,0.4]
    for i, hv in enumerate(heights):
        bh = Inches(hv * 0.9)
        rect(slide, bx + (bw + Inches(0.02)) * i, base_y - bh, bw, bh, BLUE)

def mock_threat(slide, x, y, w, h):
    browser_chrome(slide, x, y, w, h, "cspm.local/threats")
    cy = y + Inches(0.45)
    text(slide, x + Inches(0.18), cy, Inches(4), Inches(0.3),
         "Threat Detection", size=11, bold=True, color=BLACK)
    chips = [("3","CRITICAL", SEV_CRIT), ("11","HIGH", SEV_HIGH), ("18","MEDIUM", SEV_MED)]
    cx = x + Inches(0.18); cyy = cy + Inches(0.35)
    for i, (n, l, c) in enumerate(chips):
        rounded(slide, cx + (Inches(1.4) + Inches(0.05)) * i, cyy, Inches(1.4), Inches(0.45), c)
        text(slide, cx + (Inches(1.4) + Inches(0.05)) * i, cyy + Inches(0.04),
             Inches(1.4), Inches(0.2), n, size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        text(slide, cx + (Inches(1.4) + Inches(0.05)) * i, cyy + Inches(0.25),
             Inches(1.4), Inches(0.18), l, size=7, color=WHITE, align=PP_ALIGN.CENTER)
    ty = cyy + Inches(0.6)
    tabs = ["All", "Persistence", "Recon", "Exfil", "Priv Esc", "Defense"]
    tx = x + Inches(0.18)
    for i, t in enumerate(tabs):
        tw = Inches(0.85)
        bg = BLUE if i == 0 else WHITE
        tc = WHITE if i == 0 else TEXT_DARK
        ln = BLUE if i == 0 else BORDER
        rounded(slide, tx, ty, tw, Inches(0.27), bg, line=ln)
        text(slide, tx, ty + Inches(0.03), tw, Inches(0.22),
             t, size=7, bold=True, color=tc, align=PP_ALIGN.CENTER)
        tx += tw + Inches(0.05)
    rows = [
        ("CRITICAL", SEV_CRIT, "Root Account API Activity",   "AWS::Root",          "10:42"),
        ("CRITICAL", SEV_CRIT, "Mass S3 Object Deletion",     "iam::svc-backup",    "09:18"),
        ("HIGH",     SEV_HIGH, "New IAM Admin User Created",  "iam::ops-deploy",    "08:55"),
        ("HIGH",     SEV_HIGH, "Console Login Without MFA",   "iam::dev-jane",      "07:30"),
        ("MEDIUM",   SEV_MED,  "IAM Enumeration Detected",    "role::recon-bot",    "06:11"),
    ]
    ry = ty + Inches(0.4)
    for sev_l, sev_c, title, actor, time_s in rows:
        rounded(slide, x + Inches(0.18), ry, w - Inches(0.36), Inches(0.38), WHITE, line=BORDER)
        rect(slide, x + Inches(0.18), ry, Inches(0.06), Inches(0.38), sev_c)
        severity_chip(slide, x + Inches(0.3), ry + Inches(0.08), Inches(0.65), Inches(0.22), sev_l, sev_c)
        text(slide, x + Inches(1.0), ry + Inches(0.08), Inches(2.6), Inches(0.22),
             title, size=7.5, bold=True, color=BLACK)
        text(slide, x + Inches(3.7), ry + Inches(0.08), Inches(1.6), Inches(0.22),
             actor, size=7, color=BLUE)
        text(slide, x + Inches(5.3), ry + Inches(0.08), Inches(0.8), Inches(0.22),
             time_s, size=7, color=TEXT_MUTED)
        ry += Inches(0.45)

def mock_compliance(slide, x, y, w, h):
    browser_chrome(slide, x, y, w, h, "cspm.local/compliance")
    cy = y + Inches(0.45)
    text(slide, x + Inches(0.18), cy, Inches(4), Inches(0.3),
         "Compliance Score", size=11, bold=True, color=BLACK)
    fws = [("PCI-DSS", 78, SEV_HIGH),  ("SOC 2",     92, SEV_OK),
           ("ISO 27001",85,BLUE),       ("HIPAA",     71, SEV_HIGH),
           ("CIS",     88, SEV_OK),     ("GDPR",      82, BLUE)]
    fy = cy + Inches(0.4); fx = x + Inches(0.18)
    cw = (w - Inches(0.45)) / 3 - Inches(0.05)
    for i, (name, score, c) in enumerate(fws):
        col = i % 3; row = i // 3
        ix = fx + (cw + Inches(0.05)) * col
        iy = fy + (Inches(1.1) + Inches(0.05)) * row
        rounded(slide, ix, iy, cw, Inches(1.1), WHITE, line=BORDER)
        oval(slide, ix + Inches(0.12), iy + Inches(0.18), Inches(0.75), Inches(0.75), PANEL_BG)
        oval(slide, ix + Inches(0.2), iy + Inches(0.26), Inches(0.59), Inches(0.59), WHITE)
        text(slide, ix + Inches(0.12), iy + Inches(0.42), Inches(0.75), Inches(0.3),
             f"{score}%", size=11, bold=True, color=c, align=PP_ALIGN.CENTER)
        text(slide, ix + Inches(0.95), iy + Inches(0.22), cw - Inches(1.05), Inches(0.3),
             name, size=10, bold=True, color=BLACK)
        rounded(slide, ix + Inches(0.95), iy + Inches(0.58), cw - Inches(1.1), Inches(0.12), PANEL_BG)
        rounded(slide, ix + Inches(0.95), iy + Inches(0.58),
                (cw - Inches(1.1)) * (score/100), Inches(0.12), c)
        text(slide, ix + Inches(0.95), iy + Inches(0.75), cw - Inches(1.1), Inches(0.25),
             f"{score} of 100 controls passing", size=6.5, color=TEXT_MUTED)

def mock_reports(slide, x, y, w, h):
    browser_chrome(slide, x, y, w, h, "cspm.local/reports")
    cy = y + Inches(0.45)
    text(slide, x + Inches(0.18), cy, Inches(4), Inches(0.3),
         "Findings — All Clouds", size=11, bold=True, color=BLACK)
    pills = ["AWS", "S3", "CRITICAL", "Last 7d"]
    px = x + Inches(0.18); py = cy + Inches(0.35)
    for p in pills:
        rounded(slide, px, py, Inches(0.7), Inches(0.25), BLUE_LIGHT, line=BLUE_LINE)
        text(slide, px, py + Inches(0.04), Inches(0.7), Inches(0.18),
             p, size=7, color=BLUE_DEEP, bold=True, align=PP_ALIGN.CENTER)
        px += Inches(0.78)
    bx = x + w - Inches(2.0); by = py
    for label in ["PDF", "CSV", "JSON"]:
        rounded(slide, bx, by, Inches(0.55), Inches(0.25), BLUE)
        text(slide, bx, by + Inches(0.04), Inches(0.55), Inches(0.18),
             label, size=7, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        bx += Inches(0.6)
    thy = py + Inches(0.45)
    rect(slide, x + Inches(0.18), thy, w - Inches(0.36), Inches(0.3), BLUE_DEEP)
    headers = [("Sev", 0.7), ("Service", 1.0), ("Title", 2.6),
               ("Resource", 1.8), ("Status", 0.9)]
    hx = x + Inches(0.25)
    for h_label, h_w in headers:
        text(slide, hx, thy + Inches(0.06), Inches(h_w), Inches(0.2),
             h_label, size=7, bold=True, color=WHITE)
        hx += Inches(h_w)
    rows = [
        ("CRITICAL", SEV_CRIT, "S3",    "S3 Bucket Publicly Accessible",      "prod-data-logs",        "Open"),
        ("HIGH",     SEV_HIGH, "IAM",   "User Has Access Key Older Than 90d", "iam::deploy-bot",       "In-Progress"),
        ("HIGH",     SEV_HIGH, "RDS",   "RDS Instance Not Encrypted",         "rds::orders-db",        "Open"),
        ("MEDIUM",   SEV_MED,  "ECR",   "Image CVE: CVE-2024-21626 (HIGH)",   "ecr::api-service:v3.2", "Open"),
        ("LOW",      SEV_LOW,  "VPC",   "Flow Logs Disabled",                 "vpc-0a1b2c3d",          "Fixed"),
    ]
    ry = thy + Inches(0.32)
    for idx, (sev_l, sev_c, svc, title, res, status) in enumerate(rows):
        bg = WHITE if idx % 2 == 0 else PANEL_BG
        rect(slide, x + Inches(0.18), ry, w - Inches(0.36), Inches(0.28), bg)
        severity_chip(slide, x + Inches(0.25), ry + Inches(0.05), Inches(0.6), Inches(0.18), sev_l, sev_c)
        text(slide, x + Inches(0.95), ry + Inches(0.05), Inches(0.95), Inches(0.2),
             svc, size=7, bold=True, color=BLUE)
        text(slide, x + Inches(1.95), ry + Inches(0.05), Inches(2.6), Inches(0.2),
             title, size=7, color=BLACK)
        text(slide, x + Inches(4.55), ry + Inches(0.05), Inches(1.8), Inches(0.2),
             res, size=7, color=TEXT_MUTED)
        st_color = SEV_CRIT if status == "Open" else (SEV_HIGH if status == "In-Progress" else SEV_OK)
        text(slide, x + Inches(6.35), ry + Inches(0.05), Inches(0.9), Inches(0.2),
             "● " + status, size=7, bold=True, color=st_color)
        ry += Inches(0.3)

def mock_drift(slide, x, y, w, h):
    browser_chrome(slide, x, y, w, h, "cspm.local/baselines")
    cy = y + Inches(0.45)
    text(slide, x + Inches(0.18), cy, Inches(4), Inches(0.3),
         "Baseline Drift — v1.4 vs current", size=11, bold=True, color=BLACK)
    col_w = (w - Inches(0.5)) / 2
    cy2 = cy + Inches(0.4)
    rounded(slide, x + Inches(0.18), cy2, col_w, Inches(2.4), WHITE, line=BORDER)
    rect(slide, x + Inches(0.18), cy2, col_w, Inches(0.3), BLUE_DEEP)
    text(slide, x + Inches(0.18), cy2 + Inches(0.05), col_w, Inches(0.2),
         "Baseline v1.4 (snapshot 2026-04-01)", size=8, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    rx = x + Inches(0.18) + col_w + Inches(0.14)
    rounded(slide, rx, cy2, col_w, Inches(2.4), WHITE, line=BORDER)
    rect(slide, rx, cy2, col_w, Inches(0.3), SEV_CRIT)
    text(slide, rx, cy2 + Inches(0.05), col_w, Inches(0.2),
         "Current (2026-05-13) — 7 drifts", size=8, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    diffs = [
        ("S3 bucket policy",   "Block public ACL",    "Public READ added",   SEV_CRIT),
        ("SG sg-0abc",         "0 inbound from any",  "0.0.0.0/0:22 added",  SEV_HIGH),
        ("IAM role admin",     "12 attached policies","15 (FullAccess added)", SEV_HIGH),
        ("RDS encryption",     "Enabled (AES-256)",   "Disabled",            SEV_HIGH),
        ("CloudTrail logging", "Multi-region ON",     "Trail stopped",       SEV_CRIT),
    ]
    ry = cy2 + Inches(0.34)
    for idx, (label, before, after, sev) in enumerate(diffs):
        bg = PANEL_BG if idx % 2 == 0 else WHITE
        rect(slide, x + Inches(0.18), ry, col_w, Inches(0.38), bg)
        rect(slide, rx, ry, col_w, Inches(0.38), bg)
        text(slide, x + Inches(0.25), ry + Inches(0.04), col_w - Inches(0.1), Inches(0.18),
             label, size=7, bold=True, color=BLACK)
        text(slide, x + Inches(0.25), ry + Inches(0.2), col_w - Inches(0.1), Inches(0.18),
             before, size=7, color=TEXT_DARK)
        text(slide, rx + Inches(0.07), ry + Inches(0.04), col_w - Inches(0.1), Inches(0.18),
             label, size=7, bold=True, color=BLACK)
        rect(slide, rx + Inches(0.07), ry + Inches(0.2), Inches(0.08), Inches(0.18), sev)
        text(slide, rx + Inches(0.2), ry + Inches(0.2), col_w - Inches(0.25), Inches(0.18),
             after, size=7, bold=True, color=sev)
        ry += Inches(0.4)

# ============================================================
# SLIDE 1 — TITLE (white)
# ============================================================
s = add_blank()
# left accent panel
rect(s, 0, 0, Inches(0.4), SH, BLUE)
# subtle blue band
rect(s, Inches(0.4), Inches(0), SW - Inches(0.4), Inches(0.08), BLUE)
# top label
text(s, Inches(0.9), Inches(1.8), Inches(11), Inches(0.4),
     "MULTI-CLOUD CSPM PLATFORM", size=13, bold=True, color=BLUE)
text(s, Inches(0.9), Inches(2.2), Inches(11), Inches(1.8),
     "Unified Cloud Security\nin One Platform.",
     size=46, bold=True, color=BLACK)
# accent line
rect(s, Inches(0.9), Inches(4.15), Inches(1.2), Inches(0.07), BLUE)
text(s, Inches(0.9), Inches(4.35), Inches(11), Inches(0.5),
     "AWS · Azure · GCP  —  Scan, detect, comply, remediate.",
     size=16, color=TEXT_DARK)
text(s, Inches(0.9), Inches(4.9), Inches(11), Inches(0.5),
     "70+ scanners · 6 frameworks · 31 UI pages · self-hosted.",
     size=14, color=BLUE_DEEP)
# CTA pill
rounded(s, Inches(0.9), Inches(5.7), Inches(3.5), Inches(0.5), BLUE)
text(s, Inches(0.9), Inches(5.78), Inches(3.5), Inches(0.4),
     "FEATURE SHOWCASE — SHORT EDITION",
     size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
# right hero box (subtle light blue panel as visual)
rounded(s, Inches(8.6), Inches(1.8), Inches(4.2), Inches(4.6), PANEL_BG, line=BLUE_LINE)
# stat callouts inside
stats = [("70+", "Scanners"), ("3", "Clouds"), ("6", "Frameworks"), ("31", "UI Pages")]
sy = Inches(2.05)
for n, l in stats:
    text(s, Inches(8.8), sy, Inches(2), Inches(0.6),
         n, size=30, bold=True, color=BLUE)
    text(s, Inches(11.0), sy + Inches(0.18), Inches(1.8), Inches(0.4),
         l, size=12, color=TEXT_DARK)
    sy += Inches(1.05)

# ============================================================
# SLIDE 2 — PLATFORM AT A GLANCE
# ============================================================
s = add_blank()
page_header(s, "Platform at a Glance", "What you get out of the box")

stats = [
    ("70+", "Cloud Scanners"),
    ("3",   "Cloud Providers"),
    ("31",  "UI Pages"),
    ("6",   "Compliance Frameworks"),
    ("7",   "Async Workers"),
    ("41",  "Data Models"),
]
y = Inches(1.35); w = Inches(1.95); h = Inches(1.4); gap = Inches(0.12)
for i, (n, l) in enumerate(stats):
    x = Inches(0.55) + (w + gap) * i
    rounded(s, x, y, w, h, WHITE, line=BLUE_LINE)
    rect(s, x, y, w, Inches(0.06), BLUE)
    text(s, x, y + Inches(0.2), w, Inches(0.7),
         n, size=32, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    text(s, x, y + Inches(0.95), w, Inches(0.4),
         l, size=10, color=TEXT_DARK, align=PP_ALIGN.CENTER)

pillars = [
    ("Multi-Cloud, One UX", "AWS · Azure · GCP scanners share one engine, schema, dashboard."),
    ("Detect → Decide → Fix", "Findings flow into anomaly, approvals, auto-remediation."),
    ("Compliance Built-in", "PCI · SOC2 · ISO · HIPAA · CIS · GDPR — mapped automatically."),
    ("Self-Hosted by Design", "Postgres + Redis. Runs on one VM. Your data stays yours."),
]
y2 = Inches(3.1); cw = Inches(3.05); ch = Inches(1.55); gap = Inches(0.12)
for i, (t, b) in enumerate(pillars):
    x = Inches(0.55) + (cw + gap) * i
    rounded(s, x, y2, cw, ch, WHITE, line=BLUE_LINE)
    rect(s, x, y2, Inches(0.08), ch, BLUE)
    text(s, x + Inches(0.2), y2 + Inches(0.1), cw - Inches(0.25), Inches(0.4),
         t, size=13, bold=True, color=BLACK)
    text(s, x + Inches(0.2), y2 + Inches(0.5), cw - Inches(0.25), ch - Inches(0.55),
         b, size=10.5, color=TEXT_DARK)

# dashboard mock
mock_dashboard(s, Inches(0.55), Inches(4.95), Inches(12.3), Inches(2.25))
text(s, Inches(0.55), SH - Inches(0.3), Inches(12), Inches(0.22),
     "↑ Dashboard — posture score, finding counts, 30-day trend across all clouds",
     size=8, color=TEXT_MUTED)

# ============================================================
# SLIDE 3 — CLOUD COVERAGE
# ============================================================
s = add_blank()
page_header(s, "Cloud Coverage", "70+ purpose-built scanners across AWS, Azure & GCP")

clouds = [
    ("AWS",   "25+",
     ["CloudTrail · IAM · S3 · EC2 · RDS",
      "KMS · Secrets Manager · CloudWatch",
      "VPC · Lambda · ECR · EBS · ELB",
      "DynamoDB · ElastiCache · ACM",
      "SNS · SQS · Redshift · ECS · WAF",
      "CloudFront · API Gateway · SSM",
      "Threat Detection (GuardDuty-eq)"]),
    ("Azure", "32",
     ["Entra ID · IAM · Key Vault · NSG",
      "Storage · Cosmos · SQL · Postgres · MySQL",
      "VM · AKS · ACR · App Service · Functions",
      "Container Apps · App Gateway · APIM",
      "Service Bus · Event Hub · Event Grid · IoT",
      "Cognitive · AML · ADF · Synapse · Search",
      "Log Analytics · Backup · Automation · Threat"]),
    ("GCP",   "13",
     ["IAM · KMS · Secret Manager",
      "Storage · BigQuery · Cloud SQL",
      "Compute Engine · GKE · Cloud Run",
      "Cloud Functions · Pub/Sub",
      "Artifact Registry · Cloud Logging",
      "",
      "Anomaly worker + CIS GCP benchmark"]),
]
y = Inches(1.35); cw = Inches(4.1); ch = Inches(4.4); gap = Inches(0.1)
for i, (name, count, items) in enumerate(clouds):
    x = Inches(0.45) + (cw + gap) * i
    rounded(s, x, y, cw, ch, WHITE, line=BLUE_LINE)
    rect(s, x, y, cw, Inches(0.95), BLUE)
    text(s, x, y + Inches(0.12), cw, Inches(0.4),
         name, size=24, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    text(s, x, y + Inches(0.55), cw, Inches(0.3),
         f"{count} scanners", size=11, color=WHITE, align=PP_ALIGN.CENTER)
    yy = y + Inches(1.15)
    for line in items:
        if not line.strip(): yy += Inches(0.15); continue
        text(s, x + Inches(0.2), yy, cw - Inches(0.3), Inches(0.32),
             "▸  " + line, size=10, color=TEXT_DARK)
        yy += Inches(0.42)

# bottom band
rounded(s, Inches(0.45), Inches(5.95), Inches(12.5), Inches(1.05), PANEL_BG, line=BLUE_LINE)
text(s, Inches(0.65), Inches(6.05), Inches(12), Inches(0.35),
     "One Engine, One Schema, One UI", size=14, bold=True, color=BLUE_DEEP)
text(s, Inches(0.65), Inches(6.4), Inches(12), Inches(0.6),
     "All scanners share fingerprint logic, severity model, finding shape & compliance mapping. "
     "Add a new cloud or service — dashboard, reports, alerts and approvals pick it up automatically.",
     size=11, color=TEXT_DARK)

# ============================================================
# SLIDE 4 — THREAT DETECTION
# ============================================================
s = add_blank()
page_header(s, "Threat Detection — GuardDuty Without the License",
            "Active threats from CloudTrail · IAM · resource APIs — no add-on needed")

left_w = Inches(5.5)
cats = [
    ("Defense Evasion",      "CloudTrail off · Config recorder stopped"),
    ("Unauthorized Access",  "Root API · login no-MFA · brute force"),
    ("Persistence",          "New admin · new keys · cross-account"),
    ("Privilege Escalation", "iam:* attached · admin role assumed"),
    ("Reconnaissance",       "IAM enum · mass Describe/List"),
    ("Data Exfiltration",    "Mass S3 delete · snapshot made public"),
]
y = Inches(1.35)
for i, (t, b) in enumerate(cats):
    iy = y + Inches(0.6) * i
    rounded(s, Inches(0.45), iy, left_w, Inches(0.55), WHITE, line=BLUE_LINE)
    rect(s, Inches(0.45), iy, Inches(0.1), Inches(0.55), BLUE)
    text(s, Inches(0.7), iy + Inches(0.07), left_w - Inches(0.3), Inches(0.22),
         t, size=11, bold=True, color=BLACK)
    text(s, Inches(0.7), iy + Inches(0.29), left_w - Inches(0.3), Inches(0.25),
         b, size=9, color=TEXT_DARK)

# right: threat mock
mock_threat(s, Inches(6.2), Inches(1.35), Inches(6.75), Inches(5.0))
text(s, Inches(6.2), Inches(6.45), Inches(6.75), Inches(0.22),
     "↑ /threats page — severity tally, category tabs, actor & timestamp",
     size=8, color=TEXT_MUTED)

# bottom strip
rounded(s, Inches(0.45), Inches(4.95), left_w + Inches(0.05), Inches(0.7),
        BLUE_LIGHT, line=BLUE_LINE)
text(s, Inches(0.6), Inches(5.0), Inches(5.5), Inches(0.25),
     "How it works", size=10, bold=True, color=BLUE_DEEP)
text(s, Inches(0.6), Inches(5.25), Inches(5.5), Inches(0.4),
     "Targeted CloudTrail LookupEvents · rate-limited 2 rps · severity-windowed (24h/7d) · per-account scan",
     size=8.5, color=TEXT_DARK)

# ============================================================
# SLIDE 5 — COMPLIANCE
# ============================================================
s = add_blank()
page_header(s, "Compliance — PCI · SOC 2 · ISO 27001 · HIPAA · CIS · GDPR",
            "Continuous score · evidence-backed · auditor-ready")

# left: 6 framework badges
fws = [
    ("PCI-DSS",   "Card-holder data"),
    ("SOC 2",     "Trust services"),
    ("ISO 27001", "ISMS"),
    ("HIPAA",     "Health info"),
    ("CIS",       "Cloud benchmark"),
    ("GDPR",      "EU data privacy"),
]
y = Inches(1.35); cw = Inches(2.7); ch = Inches(1.05); gap = Inches(0.1)
for i, (t, b) in enumerate(fws):
    col = i % 2; row = i // 2
    x = Inches(0.45) + (cw + gap) * col
    iy = y + (ch + gap) * row
    rounded(s, x, iy, cw, ch, WHITE, line=BLUE_LINE)
    rect(s, x, iy, cw, Inches(0.42), BLUE)
    text(s, x, iy + Inches(0.06), cw, Inches(0.32),
         t, size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    text(s, x, iy + Inches(0.5), cw, Inches(0.5),
         b, size=10, color=TEXT_DARK, align=PP_ALIGN.CENTER)

# right: compliance mock
mock_compliance(s, Inches(6.4), Inches(1.35), Inches(6.55), Inches(3.65))
text(s, Inches(6.4), Inches(5.1), Inches(6.55), Inches(0.22),
     "↑ /compliance — live score per framework, drill into failing controls",
     size=8, color=TEXT_MUTED)

# bottom capabilities
caps = [
    ("Posture Score Engine",  "Weighted scoring, trend history, per-account."),
    ("Evidence Collected",    "Every finding emits an evidence record (90-day TTL)."),
    ("Control Mapping",       "Each finding linked to one or more controls."),
    ("Exportable Reports",    "Per-framework PDF / HTML / CSV on demand."),
]
y2 = Inches(5.5); cw2 = Inches(3.05); ch2 = Inches(1.55); gap2 = Inches(0.12)
for i, (t, b) in enumerate(caps):
    x = Inches(0.45) + (cw2 + gap2) * i
    rounded(s, x, y2, cw2, ch2, WHITE, line=BLUE_LINE)
    rect(s, x, y2, Inches(0.08), ch2, BLUE)
    text(s, x + Inches(0.2), y2 + Inches(0.12), cw2 - Inches(0.25), Inches(0.4),
         t, size=11, bold=True, color=BLACK)
    text(s, x + Inches(0.2), y2 + Inches(0.5), cw2 - Inches(0.25), ch2 - Inches(0.55),
         b, size=10, color=TEXT_DARK)

# ============================================================
# SLIDE 6 — DRIFT, BASELINES & CONFIG SYNC
# ============================================================
s = add_blank()
page_header(s, "Drift, Baselines & Config Sync",
            "Know exactly what changed, by whom, against which baseline")

cols = [
    ("Configuration Baselines",
     ["Capture known-good state per cloud account",
      "Versioned snapshots (BaselineVersion model)",
      "Per-resource fingerprint diffing",
      "Compare any two versions side-by-side"]),
    ("Config Change Monitor",
     ["Streaming detection of resource changes",
      "Classified: security-impact vs cosmetic",
      "Who-what-when retained per change",
      "Replayable timeline per resource"]),
    ("Freeze Windows",
     ["Block remediation in release windows",
      "Per-account / per-cloud / global",
      "Audit trail of who declared the freeze",
      "Live banner in dashboard"]),
]
y = Inches(1.35); cw = Inches(4.1); ch = Inches(2.55); gap = Inches(0.1)
for i, (t, items) in enumerate(cols):
    x = Inches(0.45) + (cw + gap) * i
    rounded(s, x, y, cw, ch, WHITE, line=BLUE_LINE)
    rect(s, x, y, cw, Inches(0.5), BLUE)
    text(s, x, y + Inches(0.1), cw, Inches(0.32),
         t, size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    yy = y + Inches(0.65)
    for line in items:
        text(s, x + Inches(0.2), yy, cw - Inches(0.3), Inches(0.4),
             "▸  " + line, size=10, color=TEXT_DARK)
        yy += Inches(0.4)

# drift mock
mock_drift(s, Inches(0.45), Inches(4.1), Inches(12.5), Inches(2.95))
text(s, Inches(0.45), SH - Inches(0.3), Inches(12), Inches(0.22),
     "↑ Baseline Drift — side-by-side compare with severity-tagged deltas",
     size=8, color=TEXT_MUTED)

# ============================================================
# SLIDE 7 — REPORTING, ALERTS & INTEGRATIONS
# ============================================================
s = add_blank()
page_header(s, "Reporting, Alerts & Integrations",
            "Right finding · right team · right time")

# 3 panels
rounded(s, Inches(0.45), Inches(1.35), Inches(4.1), Inches(2.75), WHITE, line=BLUE_LINE)
rect(s, Inches(0.45), Inches(1.35), Inches(4.1), Inches(0.5), BLUE)
text(s, Inches(0.45), Inches(1.43), Inches(4.1), Inches(0.4),
     "Reporting", size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
rb = ["Per-account / per-cloud / per-framework",
      "Formats: PDF · HTML · CSV · JSON",
      "Scheduled (cron) + on-demand",
      "ReportSchedule + ReportRun audit trail",
      "Posture history exported for board"]
yy = Inches(1.95)
for line in rb:
    text(s, Inches(0.65), yy, Inches(3.85), Inches(0.38),
         "▸  " + line, size=10, color=TEXT_DARK); yy += Inches(0.4)

rounded(s, Inches(4.65), Inches(1.35), Inches(4.1), Inches(2.75), WHITE, line=BLUE_LINE)
rect(s, Inches(4.65), Inches(1.35), Inches(4.1), Inches(0.5), BLUE_DEEP)
text(s, Inches(4.65), Inches(1.43), Inches(4.1), Inches(0.4),
     "Alerts", size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
al = ["Slack channel routing by severity",
      "Email digest or per-event",
      "AlertLog — dedup + throttle aware",
      "Freeze-window aware (no noise)",
      "Quiet-hours configurable"]
yy = Inches(1.95)
for line in al:
    text(s, Inches(4.85), yy, Inches(3.85), Inches(0.38),
         "▸  " + line, size=10, color=TEXT_DARK); yy += Inches(0.4)

rounded(s, Inches(8.85), Inches(1.35), Inches(4.1), Inches(2.75), WHITE, line=BLUE_LINE)
rect(s, Inches(8.85), Inches(1.35), Inches(4.1), Inches(0.5), BLUE)
text(s, Inches(8.85), Inches(1.43), Inches(4.1), Inches(0.4),
     "Integrations", size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
ig = ["Jira / ServiceNow ticketing",
      "Splunk / DataDog / SIEM webhooks",
      "Slack · Email · generic webhook",
      "Per-finding routing rules",
      "IntegrationLog for audit"]
yy = Inches(1.95)
for line in ig:
    text(s, Inches(9.05), yy, Inches(3.85), Inches(0.38),
         "▸  " + line, size=10, color=TEXT_DARK); yy += Inches(0.4)

# reports mock
mock_reports(s, Inches(0.45), Inches(4.25), Inches(12.5), Inches(2.85))
text(s, Inches(0.45), SH - Inches(0.3), Inches(12), Inches(0.22),
     "↑ Findings/Reports — filters, severity-aware rows, PDF/CSV/JSON export, inline status",
     size=8, color=TEXT_MUTED)

# ============================================================
# SLIDE 8 — SUMMARY (white)
# ============================================================
s = add_blank()
rect(s, 0, 0, Inches(0.4), SH, BLUE)
rect(s, Inches(0.4), 0, SW - Inches(0.4), Inches(0.08), BLUE)
text(s, Inches(0.9), Inches(0.9), Inches(11), Inches(0.4),
     "SUMMARY", size=13, bold=True, color=BLUE)
text(s, Inches(0.9), Inches(1.3), Inches(11), Inches(1.4),
     "One platform.  Three clouds.\nZero per-asset pricing.",
     size=38, bold=True, color=BLACK)
rect(s, Inches(0.9), Inches(3.1), Inches(1.2), Inches(0.07), BLUE)

caps = [
    ("70+",  "Scanners"),
    ("6",    "Frameworks"),
    ("3",    "Clouds"),
    ("31",   "UI Pages"),
]
y = Inches(3.5); w = Inches(1.8); h = Inches(1.2); gap = Inches(0.15)
for i, (n, l) in enumerate(caps):
    x = Inches(0.9) + (w + gap) * i
    rounded(s, x, y, w, h, WHITE, line=BLUE_LINE)
    rect(s, x, y, w, Inches(0.06), BLUE)
    text(s, x, y + Inches(0.18), w, Inches(0.6),
         n, size=28, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    text(s, x, y + Inches(0.78), w, Inches(0.4),
         l, size=10, color=TEXT_DARK, align=PP_ALIGN.CENTER)

yv = Inches(5.05)
vals = [
    "▸  Replaces Wiz / Orca / Prisma + GuardDuty + Inspector + Macie",
    "▸  Detect → Approve → Auto-remediate, end to end",
    "▸  Self-hosted — your data, your VPC, predictable cost",
    "▸  Up and running in an afternoon (docker-compose up)",
]
for line in vals:
    text(s, Inches(0.9), yv, Inches(11.5), Inches(0.32),
         line, size=13, color=TEXT_DARK); yv += Inches(0.35)

# CTAs
rounded(s, Inches(0.9), SH - Inches(0.85), Inches(2.6), Inches(0.5), BLUE)
text(s, Inches(0.9), SH - Inches(0.77), Inches(2.6), Inches(0.4),
     "REQUEST DEMO", size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
rounded(s, Inches(3.65), SH - Inches(0.85), Inches(2.6), Inches(0.5), WHITE, line=BLUE)
text(s, Inches(3.65), SH - Inches(0.77), Inches(2.6), Inches(0.4),
     "VIEW DOCS", size=11, bold=True, color=BLUE, align=PP_ALIGN.CENTER)

# ---------- save ----------
out = "/mnt/d/Data/Linux-Data/IML/vul-scanner/aws-scanner/docs/CSPM_Platform_Showcase_Short_White.pptx"
prs.save(out)
print("Saved:", out)
print("Slides:", len(prs.slides.__iter__.__self__._sldIdLst))
