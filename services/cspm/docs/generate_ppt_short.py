"""Short showcase PPT — 8 slides with mocked UI screenshots."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---------- palette ----------
NAVY      = RGBColor(0x0B, 0x1F, 0x3A)
DEEP_BLUE = RGBColor(0x13, 0x2C, 0x55)
ACCENT    = RGBColor(0x2D, 0x9C, 0xDB)
TEAL      = RGBColor(0x1A, 0xBC, 0x9C)
ORANGE    = RGBColor(0xE6, 0x7E, 0x22)
RED       = RGBColor(0xC0, 0x39, 0x2B)
YELLOW    = RGBColor(0xF1, 0xC4, 0x0F)
GRAY_BG   = RGBColor(0xF4, 0xF6, 0xF8)
GRAY_DARK = RGBColor(0x34, 0x49, 0x5E)
GRAY_MID  = RGBColor(0x7B, 0x8A, 0x9D)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRY = RGBColor(0xDC, 0xE0, 0xE6)
SOFT_GREEN= RGBColor(0x27, 0xAE, 0x60)

# severity colors per memory
SEV_CRIT  = RGBColor(0xDC, 0x26, 0x26)
SEV_HIGH  = RGBColor(0xEA, 0x58, 0x0C)
SEV_MED   = RGBColor(0xEA, 0xB3, 0x08)
SEV_LOW   = RGBColor(0x3B, 0x82, 0xF6)
SEV_INFO  = RGBColor(0x6B, 0x72, 0x80)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height

# ---------- helpers ----------
def add_blank():
    return prs.slides.add_slide(prs.slide_layouts[6])

def rect(slide, x, y, w, h, fill, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb = line
    shp.shadow.inherit = False
    return shp

def rounded(slide, x, y, w, h, fill, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb = line
    shp.shadow.inherit = False
    try:
        shp.adjustments[0] = 0.1
    except Exception: pass
    return shp

def oval(slide, x, y, w, h, fill, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb = line
    shp.shadow.inherit = False
    return shp

def text(slide, x, y, w, h, body, *, size=14, bold=False, color=GRAY_DARK,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Calibri"):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    lines = body.split("\n") if isinstance(body, str) else body
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = font
    return tb

def page_header(s, title, subtitle=None):
    rect(s, 0, 0, SW, Inches(0.85), NAVY)
    rect(s, 0, Inches(0.85), Inches(0.18), SH - Inches(0.85), ACCENT)
    text(s, Inches(0.45), Inches(0.13), Inches(12.5), Inches(0.5),
         title, size=24, bold=True, color=WHITE)
    if subtitle:
        text(s, Inches(0.45), Inches(0.5), Inches(12.5), Inches(0.35),
             subtitle, size=11, color=RGBColor(0xCB, 0xD7, 0xE3))
    text(s, Inches(0.4), SH - Inches(0.3), Inches(12), Inches(0.22),
         "Multi-Cloud CSPM Platform  |  AWS • Azure • GCP",
         size=8, color=GRAY_MID)

# ---------- mock UI builders ----------
def browser_chrome(slide, x, y, w, h, url):
    rounded(slide, x, y, w, h, WHITE, line=LIGHT_GRY)
    rect(slide, x + Inches(0.04), y + Inches(0.04), w - Inches(0.08), Inches(0.32), GRAY_BG)
    # window dots
    for i, c in enumerate([RGBColor(0xEF,0x46,0x46), RGBColor(0xF5,0x9E,0x0B), RGBColor(0x10,0xB9,0x81)]):
        oval(slide, x + Inches(0.12) + Inches(0.18) * i, y + Inches(0.12),
             Inches(0.12), Inches(0.12), c)
    # url bar
    rounded(slide, x + Inches(0.85), y + Inches(0.08), w - Inches(1.0), Inches(0.24),
            WHITE, line=LIGHT_GRY)
    text(slide, x + Inches(0.95), y + Inches(0.1), w - Inches(1.2), Inches(0.2),
         url, size=8, color=GRAY_MID)
    # content area starts at y + 0.4

def severity_chip(slide, x, y, w, h, label, color):
    rounded(slide, x, y, w, h, color)
    text(slide, x, y + Inches(0.03), w, h - Inches(0.03),
         label, size=7, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

def mock_dashboard(slide, x, y, w, h):
    """Mock of Dashboard page."""
    browser_chrome(slide, x, y, w, h, "cspm.local/dashboard")
    cy = y + Inches(0.45)
    # title
    text(slide, x + Inches(0.18), cy, Inches(3), Inches(0.3),
         "Security Posture", size=11, bold=True, color=NAVY)
    text(slide, x + Inches(0.18), cy + Inches(0.25), Inches(3), Inches(0.2),
         "Across 3 clouds, 12 accounts", size=7, color=GRAY_MID)
    # KPI cards
    kpis = [("87", "Posture", SOFT_GREEN), ("142", "Findings", ORANGE),
            ("9", "Critical", SEV_CRIT), ("31", "Resources", ACCENT)]
    kx = x + Inches(0.18); ky = cy + Inches(0.6)
    kw = (w - Inches(0.45)) / 4 - Inches(0.05)
    for i, (n, l, c) in enumerate(kpis):
        rounded(slide, kx + (kw + Inches(0.05)) * i, ky, kw, Inches(0.7), WHITE, line=LIGHT_GRY)
        rect(slide, kx + (kw + Inches(0.05)) * i, ky, Inches(0.05), Inches(0.7), c)
        text(slide, kx + (kw + Inches(0.05)) * i + Inches(0.12), ky + Inches(0.05),
             kw - Inches(0.12), Inches(0.32), n, size=15, bold=True, color=NAVY)
        text(slide, kx + (kw + Inches(0.05)) * i + Inches(0.12), ky + Inches(0.42),
             kw - Inches(0.12), Inches(0.25), l, size=7, color=GRAY_MID)
    # chart area
    cy2 = ky + Inches(0.85)
    rounded(slide, x + Inches(0.18), cy2, w - Inches(0.36), h - (cy2 - y) - Inches(0.18),
            GRAY_BG)
    text(slide, x + Inches(0.3), cy2 + Inches(0.08), Inches(3), Inches(0.2),
         "Findings trend (30d)", size=8, bold=True, color=NAVY)
    # bar chart
    base_y = cy2 + h - (cy2 - y) - Inches(0.3)
    bx = x + Inches(0.3)
    bw = (w - Inches(0.6)) / 14
    heights = [0.25,0.35,0.3,0.5,0.4,0.45,0.6,0.5,0.7,0.55,0.45,0.6,0.5,0.4]
    for i, hv in enumerate(heights):
        bh = Inches(hv * 0.9)
        rect(slide, bx + (bw + Inches(0.02)) * i, base_y - bh, bw, bh, ACCENT)

def mock_threat(slide, x, y, w, h):
    """Mock of Threat Detection page."""
    browser_chrome(slide, x, y, w, h, "cspm.local/threats")
    cy = y + Inches(0.45)
    text(slide, x + Inches(0.18), cy, Inches(4), Inches(0.3),
         "Threat Detection", size=11, bold=True, color=NAVY)
    # summary chips
    chips = [("3", "CRITICAL", SEV_CRIT), ("11", "HIGH", SEV_HIGH), ("18", "MEDIUM", SEV_MED)]
    cx = x + Inches(0.18); cyy = cy + Inches(0.35)
    for i, (n, l, c) in enumerate(chips):
        rounded(slide, cx + (Inches(1.4) + Inches(0.05)) * i, cyy, Inches(1.4), Inches(0.45), c)
        text(slide, cx + (Inches(1.4) + Inches(0.05)) * i, cyy + Inches(0.04),
             Inches(1.4), Inches(0.2), n, size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        text(slide, cx + (Inches(1.4) + Inches(0.05)) * i, cyy + Inches(0.25),
             Inches(1.4), Inches(0.18), l, size=7, color=WHITE, align=PP_ALIGN.CENTER)
    # tab row
    ty = cyy + Inches(0.6)
    tabs = ["All", "Persistence", "Recon", "Exfil", "Priv Esc", "Defense"]
    tx = x + Inches(0.18)
    for i, t in enumerate(tabs):
        tw = Inches(0.85)
        bg = ACCENT if i == 0 else WHITE
        tc = WHITE if i == 0 else GRAY_DARK
        rounded(slide, tx, ty, tw, Inches(0.27), bg, line=LIGHT_GRY)
        text(slide, tx, ty + Inches(0.03), tw, Inches(0.22),
             t, size=7, bold=True, color=tc, align=PP_ALIGN.CENTER)
        tx += tw + Inches(0.05)
    # threat rows
    rows = [
        ("CRITICAL", SEV_CRIT, "Root Account API Activity",   "AWS::Root",          "10:42"),
        ("CRITICAL", SEV_CRIT, "Mass S3 Object Deletion",     "iam::svc-backup",    "09:18"),
        ("HIGH",     SEV_HIGH, "New IAM Admin User Created",  "iam::ops-deploy",    "08:55"),
        ("HIGH",     SEV_HIGH, "Console Login Without MFA",   "iam::dev-jane",      "07:30"),
        ("MEDIUM",   SEV_MED,  "IAM Enumeration Detected",    "role::recon-bot",    "06:11"),
    ]
    ry = ty + Inches(0.4)
    for sev_l, sev_c, title, actor, time_s in rows:
        rounded(slide, x + Inches(0.18), ry, w - Inches(0.36), Inches(0.38), WHITE, line=LIGHT_GRY)
        rect(slide, x + Inches(0.18), ry, Inches(0.06), Inches(0.38), sev_c)
        severity_chip(slide, x + Inches(0.3), ry + Inches(0.08), Inches(0.65), Inches(0.22), sev_l, sev_c)
        text(slide, x + Inches(1.0), ry + Inches(0.08), Inches(2.6), Inches(0.22),
             title, size=7.5, bold=True, color=NAVY)
        text(slide, x + Inches(3.7), ry + Inches(0.08), Inches(1.6), Inches(0.22),
             actor, size=7, color=GRAY_DARK)
        text(slide, x + Inches(5.3), ry + Inches(0.08), Inches(0.8), Inches(0.22),
             time_s, size=7, color=GRAY_MID)
        ry += Inches(0.45)

def mock_compliance(slide, x, y, w, h):
    """Mock of Compliance page."""
    browser_chrome(slide, x, y, w, h, "cspm.local/compliance")
    cy = y + Inches(0.45)
    text(slide, x + Inches(0.18), cy, Inches(4), Inches(0.3),
         "Compliance Score", size=11, bold=True, color=NAVY)
    # framework rings
    fws = [("PCI-DSS", 78, ORANGE), ("SOC 2", 92, SOFT_GREEN),
           ("ISO 27001", 85, TEAL), ("HIPAA", 71, ORANGE),
           ("CIS", 88, SOFT_GREEN), ("GDPR", 82, TEAL)]
    fy = cy + Inches(0.4); fx = x + Inches(0.18)
    cw = (w - Inches(0.45)) / 3 - Inches(0.05)
    for i, (name, score, c) in enumerate(fws):
        col = i % 3; row = i // 3
        ix = fx + (cw + Inches(0.05)) * col
        iy = fy + (Inches(1.1) + Inches(0.05)) * row
        rounded(slide, ix, iy, cw, Inches(1.1), WHITE, line=LIGHT_GRY)
        # circle gauge
        oval(slide, ix + Inches(0.12), iy + Inches(0.18), Inches(0.75), Inches(0.75),
             GRAY_BG)
        oval(slide, ix + Inches(0.2), iy + Inches(0.26), Inches(0.59), Inches(0.59), WHITE)
        text(slide, ix + Inches(0.12), iy + Inches(0.42), Inches(0.75), Inches(0.3),
             f"{score}%", size=11, bold=True, color=c, align=PP_ALIGN.CENTER)
        # label
        text(slide, ix + Inches(0.95), iy + Inches(0.22), cw - Inches(1.05), Inches(0.3),
             name, size=10, bold=True, color=NAVY)
        # progress bar
        rounded(slide, ix + Inches(0.95), iy + Inches(0.58), cw - Inches(1.1), Inches(0.12),
                GRAY_BG)
        rounded(slide, ix + Inches(0.95), iy + Inches(0.58),
                (cw - Inches(1.1)) * (score/100), Inches(0.12), c)
        text(slide, ix + Inches(0.95), iy + Inches(0.75), cw - Inches(1.1), Inches(0.25),
             f"{score} of 100 controls passing", size=6.5, color=GRAY_MID)

def mock_reports(slide, x, y, w, h):
    """Mock of Reports / Findings page."""
    browser_chrome(slide, x, y, w, h, "cspm.local/reports")
    cy = y + Inches(0.45)
    text(slide, x + Inches(0.18), cy, Inches(4), Inches(0.3),
         "Findings — All Clouds", size=11, bold=True, color=NAVY)
    # filter pills
    pills = ["AWS", "S3", "CRITICAL", "Last 7d"]
    px = x + Inches(0.18); py = cy + Inches(0.35)
    for p in pills:
        rounded(slide, px, py, Inches(0.7), Inches(0.25), GRAY_BG, line=LIGHT_GRY)
        text(slide, px, py + Inches(0.04), Inches(0.7), Inches(0.18),
             p, size=7, color=GRAY_DARK, align=PP_ALIGN.CENTER)
        px += Inches(0.78)
    # export buttons (right)
    bx = x + w - Inches(2.0); by = py
    for label, c in [("PDF", RED), ("CSV", TEAL), ("JSON", ACCENT)]:
        rounded(slide, bx, by, Inches(0.55), Inches(0.25), c)
        text(slide, bx, by + Inches(0.04), Inches(0.55), Inches(0.18),
             label, size=7, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        bx += Inches(0.6)
    # table header
    thy = py + Inches(0.45)
    rect(slide, x + Inches(0.18), thy, w - Inches(0.36), Inches(0.3), NAVY)
    headers = [("Sev", 0.7), ("Service", 1.0), ("Title", 2.6),
               ("Resource", 1.8), ("Status", 0.9)]
    hx = x + Inches(0.25)
    for h_label, h_w in headers:
        text(slide, hx, thy + Inches(0.06), Inches(h_w), Inches(0.2),
             h_label, size=7, bold=True, color=WHITE)
        hx += Inches(h_w)
    # rows
    rows = [
        ("CRITICAL", SEV_CRIT, "S3",    "S3 Bucket Publicly Accessible",      "prod-data-logs",        "Open"),
        ("HIGH",     SEV_HIGH, "IAM",   "User Has Access Key Older Than 90d", "iam::deploy-bot",       "In-Progress"),
        ("HIGH",     SEV_HIGH, "RDS",   "RDS Instance Not Encrypted",         "rds::orders-db",        "Open"),
        ("MEDIUM",   SEV_MED,  "ECR",   "Image CVE: CVE-2024-21626 (HIGH)",   "ecr::api-service:v3.2", "Open"),
        ("LOW",      SEV_LOW,  "VPC",   "Flow Logs Disabled",                 "vpc-0a1b2c3d",          "Fixed"),
    ]
    ry = thy + Inches(0.32)
    for sev_l, sev_c, svc, title, res, status in rows:
        bg = WHITE if rows.index((sev_l, sev_c, svc, title, res, status)) % 2 == 0 else GRAY_BG
        rect(slide, x + Inches(0.18), ry, w - Inches(0.36), Inches(0.28), bg)
        # sev chip
        severity_chip(slide, x + Inches(0.25), ry + Inches(0.05), Inches(0.6), Inches(0.18), sev_l, sev_c)
        text(slide, x + Inches(0.95), ry + Inches(0.05), Inches(0.95), Inches(0.2),
             svc, size=7, bold=True, color=NAVY)
        text(slide, x + Inches(1.95), ry + Inches(0.05), Inches(2.6), Inches(0.2),
             title, size=7, color=GRAY_DARK)
        text(slide, x + Inches(4.55), ry + Inches(0.05), Inches(1.8), Inches(0.2),
             res, size=7, color=GRAY_MID)
        st_color = SEV_CRIT if status == "Open" else (ORANGE if status == "In-Progress" else SOFT_GREEN)
        text(slide, x + Inches(6.35), ry + Inches(0.05), Inches(0.9), Inches(0.2),
             "● " + status, size=7, bold=True, color=st_color)
        ry += Inches(0.3)

def mock_drift(slide, x, y, w, h):
    """Mock of Baseline Drift / Config Changes page."""
    browser_chrome(slide, x, y, w, h, "cspm.local/baselines")
    cy = y + Inches(0.45)
    text(slide, x + Inches(0.18), cy, Inches(4), Inches(0.3),
         "Baseline Drift — v1.4 vs current", size=11, bold=True, color=NAVY)
    # two columns
    col_w = (w - Inches(0.5)) / 2
    cy2 = cy + Inches(0.4)
    # left baseline
    rounded(slide, x + Inches(0.18), cy2, col_w, Inches(2.6), WHITE, line=LIGHT_GRY)
    rect(slide, x + Inches(0.18), cy2, col_w, Inches(0.3), DEEP_BLUE)
    text(slide, x + Inches(0.18), cy2 + Inches(0.05), col_w, Inches(0.2),
         "Baseline v1.4 (snapshot 2026-04-01)", size=8, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    # right current
    rx = x + Inches(0.18) + col_w + Inches(0.14)
    rounded(slide, rx, cy2, col_w, Inches(2.6), WHITE, line=LIGHT_GRY)
    rect(slide, rx, cy2, col_w, Inches(0.3), RED)
    text(slide, rx, cy2 + Inches(0.05), col_w, Inches(0.2),
         "Current (2026-05-13) — 7 drifts", size=8, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    # drift rows
    diffs = [
        ("S3 bucket policy",   "Block public ACL",    "Public READ added",  SEV_CRIT),
        ("SG sg-0abc",         "0 inbound from any",  "0.0.0.0/0:22 added", SEV_HIGH),
        ("IAM role admin",     "12 attached policies","15 (added FullAccess)", SEV_HIGH),
        ("RDS encryption",     "Enabled (AES-256)",   "Disabled",           SEV_HIGH),
        ("CloudTrail logging", "Multi-region ON",     "Trail stopped",      SEV_CRIT),
    ]
    ry = cy2 + Inches(0.38)
    for label, before, after, sev in diffs:
        rect(slide, x + Inches(0.18), ry, col_w, Inches(0.4), GRAY_BG if diffs.index((label, before, after, sev)) % 2 == 0 else WHITE)
        rect(slide, rx, ry, col_w, Inches(0.4), GRAY_BG if diffs.index((label, before, after, sev)) % 2 == 0 else WHITE)
        text(slide, x + Inches(0.25), ry + Inches(0.05), col_w - Inches(0.1), Inches(0.18),
             label, size=7, bold=True, color=NAVY)
        text(slide, x + Inches(0.25), ry + Inches(0.21), col_w - Inches(0.1), Inches(0.18),
             before, size=7, color=GRAY_DARK)
        text(slide, rx + Inches(0.07), ry + Inches(0.05), col_w - Inches(0.1), Inches(0.18),
             label, size=7, bold=True, color=NAVY)
        rect(slide, rx + Inches(0.07), ry + Inches(0.21), Inches(0.08), Inches(0.18), sev)
        text(slide, rx + Inches(0.2), ry + Inches(0.21), col_w - Inches(0.25), Inches(0.18),
             after, size=7, bold=True, color=sev)
        ry += Inches(0.42)

# ============================================================
# SLIDE 1 — TITLE
# ============================================================
s = add_blank()
rect(s, 0, 0, SW, SH, NAVY)
acc = s.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, Inches(8.3), 0, Inches(5.0), SH)
acc.fill.solid(); acc.fill.fore_color.rgb = DEEP_BLUE; acc.line.fill.background()
rect(s, 0, Inches(3.05), Inches(0.5), Inches(0.08), ACCENT)
text(s, Inches(0.7), Inches(2.2), Inches(11), Inches(0.5),
     "MULTI-CLOUD CSPM PLATFORM", size=13, bold=True, color=ACCENT)
text(s, Inches(0.7), Inches(2.6), Inches(11), Inches(1.4),
     "Unified Cloud Security\nin One Platform.",
     size=42, bold=True, color=WHITE)
text(s, Inches(0.7), Inches(4.5), Inches(10), Inches(0.5),
     "AWS · Azure · GCP  —  Scan, detect, comply, remediate.",
     size=15, color=RGBColor(0xCB, 0xD7, 0xE3))
rounded(s, Inches(0.7), Inches(5.5), Inches(3.3), Inches(0.5), ACCENT)
text(s, Inches(0.7), Inches(5.58), Inches(3.3), Inches(0.4),
     "FEATURE SHOWCASE — SHORT EDITION",
     size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

# ============================================================
# SLIDE 2 — PLATFORM AT A GLANCE
# ============================================================
s = add_blank()
page_header(s, "Platform at a Glance", "What you get out of the box")
stats = [
    ("70+", "Cloud Scanners",   ACCENT),
    ("3",   "Cloud Providers",  TEAL),
    ("31",  "UI Pages",         DEEP_BLUE),
    ("6",   "Compliance Frameworks", ORANGE),
    ("7",   "Async Workers",    RED),
    ("41",  "Data Models",      GRAY_DARK),
]
y = Inches(1.15); w = Inches(1.95); h = Inches(1.4); gap = Inches(0.12)
for i, (n, l, c) in enumerate(stats):
    x = Inches(0.55) + (w + gap) * i
    rounded(s, x, y, w, h, c)
    text(s, x, y + Inches(0.15), w, Inches(0.7),
         n, size=34, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    text(s, x, y + Inches(0.85), w, Inches(0.45),
         l, size=10, color=WHITE, align=PP_ALIGN.CENTER)

# four pillars
pillars = [
    ("Multi-Cloud, One UX", "AWS · Azure · GCP scanners share one engine, schema, dashboard.", ACCENT),
    ("Detect → Decide → Fix", "Findings flow into anomaly, approvals, auto-remediation.", TEAL),
    ("Compliance Built-in", "PCI · SOC2 · ISO · HIPAA · CIS · GDPR — mapped automatically.", ORANGE),
    ("Self-Hosted by Design", "Postgres + Redis. Runs on one VM. Your data stays yours.", DEEP_BLUE),
]
y2 = Inches(2.9); cw = Inches(3.05); ch = Inches(1.6); gap = Inches(0.12)
for i, (t, b, c) in enumerate(pillars):
    x = Inches(0.55) + (cw + gap) * i
    rounded(s, x, y2, cw, ch, WHITE, line=LIGHT_GRY)
    rect(s, x, y2, Inches(0.08), ch, c)
    text(s, x + Inches(0.2), y2 + Inches(0.1), cw - Inches(0.25), Inches(0.4),
         t, size=13, bold=True, color=NAVY)
    text(s, x + Inches(0.2), y2 + Inches(0.5), cw - Inches(0.25), ch - Inches(0.55),
         b, size=10.5, color=GRAY_DARK)

# bottom dashboard mock screenshot
mock_dashboard(s, Inches(0.55), Inches(4.8), Inches(12.3), Inches(2.4))
text(s, Inches(0.55), SH - Inches(0.35), Inches(12), Inches(0.22),
     "↑ Dashboard view: posture score, finding counts, 30-day trend across all clouds",
     size=8, color=GRAY_MID, align=PP_ALIGN.LEFT)

# ============================================================
# SLIDE 3 — CLOUD COVERAGE
# ============================================================
s = add_blank()
page_header(s, "Cloud Coverage", "70+ purpose-built scanners across AWS, Azure & GCP")

clouds = [
    ("AWS",   "25+",
     ["CloudTrail · IAM · S3 · EC2 · RDS",
      "KMS · Secrets Manager · CloudWatch",
      "VPC · Lambda · ECR · EBS · ELB",
      "DynamoDB · ElastiCache · ACM",
      "SNS · SQS · Redshift · ECS · WAF",
      "CloudFront · API Gateway · SSM",
      "Threat Detection (GuardDuty-eq)"],
     RGBColor(0xFF, 0x99, 0x00)),
    ("Azure", "32",
     ["Entra ID · IAM · Key Vault · NSG",
      "Storage · Cosmos · SQL · Postgres · MySQL",
      "VM · AKS · ACR · App Service · Functions",
      "Container Apps · App Gateway · APIM",
      "Service Bus · Event Hub · Event Grid · IoT",
      "Cognitive · AML · ADF · Synapse · Search",
      "Log Analytics · Backup · Automation · Threat"],
     RGBColor(0x00, 0x78, 0xD4)),
    ("GCP",   "13",
     ["IAM · KMS · Secret Manager",
      "Storage · BigQuery · Cloud SQL",
      "Compute Engine · GKE · Cloud Run",
      "Cloud Functions · Pub/Sub",
      "Artifact Registry · Cloud Logging",
      "",
      "Anomaly detection worker + CIS GCP"],
     RGBColor(0x42, 0x85, 0xF4)),
]
y = Inches(1.15); cw = Inches(4.1); ch = Inches(4.5); gap = Inches(0.1)
for i, (name, count, items, c) in enumerate(clouds):
    x = Inches(0.45) + (cw + gap) * i
    rounded(s, x, y, cw, ch, WHITE, line=LIGHT_GRY)
    rect(s, x, y, cw, Inches(0.85), c)
    text(s, x, y + Inches(0.12), cw, Inches(0.4),
         name, size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    text(s, x, y + Inches(0.5), cw, Inches(0.3),
         f"{count} scanners", size=11, color=WHITE, align=PP_ALIGN.CENTER)
    yy = y + Inches(1.05)
    for line in items:
        if not line.strip(): yy += Inches(0.15); continue
        text(s, x + Inches(0.2), yy, cw - Inches(0.3), Inches(0.32),
             "▸  " + line, size=10, color=GRAY_DARK)
        yy += Inches(0.42)

# shared engine bar
rounded(s, Inches(0.45), Inches(5.85), Inches(12.5), Inches(1.15), NAVY)
text(s, Inches(0.65), Inches(5.95), Inches(12), Inches(0.4),
     "One Engine, One Schema, One UI", size=14, bold=True, color=WHITE)
text(s, Inches(0.65), Inches(6.3), Inches(12), Inches(0.7),
     "All scanners share the same fingerprint logic, severity model, finding shape, and compliance mapping. "
     "Add a new cloud or service and the dashboard, reports, alerts, and approval flow pick it up automatically.",
     size=11, color=RGBColor(0xCB, 0xD7, 0xE3))

# ============================================================
# SLIDE 4 — THREAT DETECTION
# ============================================================
s = add_blank()
page_header(s, "Threat Detection — GuardDuty Without the License",
            "Active threats from CloudTrail · IAM · resource APIs — no add-on needed")

# left: categories
left_w = Inches(5.5)
cats = [
    ("Defense Evasion",      "CloudTrail off · Config recorder stopped", RED),
    ("Unauthorized Access",  "Root API · login no-MFA · brute force",    RED),
    ("Persistence",          "New admin · new keys · cross-account",     ORANGE),
    ("Privilege Escalation", "iam:* attached · admin role assumed",      ORANGE),
    ("Reconnaissance",       "IAM enum · mass Describe/List",            YELLOW),
    ("Data Exfiltration",    "Mass S3 delete · snapshot made public",    RED),
]
y = Inches(1.15)
for i, (t, b, c) in enumerate(cats):
    iy = y + Inches(0.65) * i
    rounded(s, Inches(0.45), iy, left_w, Inches(0.58), WHITE, line=LIGHT_GRY)
    rect(s, Inches(0.45), iy, Inches(0.12), Inches(0.58), c)
    text(s, Inches(0.7), iy + Inches(0.08), left_w - Inches(0.3), Inches(0.22),
         t, size=11, bold=True, color=NAVY)
    text(s, Inches(0.7), iy + Inches(0.3), left_w - Inches(0.3), Inches(0.25),
         b, size=9, color=GRAY_DARK)

# right: mock screenshot
mock_threat(s, Inches(6.2), Inches(1.15), Inches(6.75), Inches(5.0))
text(s, Inches(6.2), Inches(6.3), Inches(6.75), Inches(0.25),
     "↑ Threat Detection page — severity tally, category tabs, actor & event time",
     size=8, color=GRAY_MID)

# how it works bottom strip
rounded(s, Inches(0.45), Inches(6.55), left_w + Inches(0.05), Inches(0.7), NAVY)
text(s, Inches(0.55), Inches(6.6), Inches(5.5), Inches(0.25),
     "How it works", size=10, bold=True, color=WHITE)
text(s, Inches(0.55), Inches(6.85), Inches(5.5), Inches(0.35),
     "Targeted CloudTrail LookupEvents · rate-limited 2 rps · severity-windowed (24h/7d) · per-account scan",
     size=8.5, color=RGBColor(0xCB, 0xD7, 0xE3))

# ============================================================
# SLIDE 5 — COMPLIANCE
# ============================================================
s = add_blank()
page_header(s, "Compliance — PCI · SOC 2 · ISO 27001 · HIPAA · CIS · GDPR",
            "Continuous score, evidence-backed, auditor-ready")

# left: framework badges
fws = [
    ("PCI-DSS",   "Card-holder data",     RED),
    ("SOC 2",     "Trust services",       ACCENT),
    ("ISO 27001", "ISMS",                 TEAL),
    ("HIPAA",     "Health info",          ORANGE),
    ("CIS",       "Cloud benchmark",      DEEP_BLUE),
    ("GDPR",      "EU data privacy",      GRAY_DARK),
]
y = Inches(1.15); cw = Inches(2.7); ch = Inches(1.1); gap = Inches(0.1)
for i, (t, b, c) in enumerate(fws):
    col = i % 2; row = i // 2
    x = Inches(0.45) + (cw + gap) * col
    iy = y + (ch + gap) * row
    rounded(s, x, iy, cw, ch, WHITE, line=LIGHT_GRY)
    rect(s, x, iy, cw, Inches(0.42), c)
    text(s, x, iy + Inches(0.06), cw, Inches(0.32),
         t, size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    text(s, x, iy + Inches(0.5), cw, Inches(0.55),
         b, size=10, color=GRAY_DARK, align=PP_ALIGN.CENTER)

# right: mock screenshot
mock_compliance(s, Inches(6.4), Inches(1.15), Inches(6.55), Inches(3.8))
text(s, Inches(6.4), Inches(5.05), Inches(6.55), Inches(0.22),
     "↑ Compliance page — live score per framework, drill into failing controls",
     size=8, color=GRAY_MID)

# bottom: capabilities
caps = [
    ("Posture Score Engine",     "Weighted scoring, trend history, per-account."),
    ("Evidence Auto-Collected",  "Every finding emits an evidence record (90-day TTL)."),
    ("Control Mapping",          "Each finding linked to one or more controls."),
    ("Exportable Reports",       "Per-framework PDF / HTML / CSV on demand."),
]
y2 = Inches(5.5); cw2 = Inches(3.05); ch2 = Inches(1.55); gap2 = Inches(0.12)
for i, (t, b) in enumerate(caps):
    x = Inches(0.45) + (cw2 + gap2) * i
    rounded(s, x, y2, cw2, ch2, GRAY_BG, line=LIGHT_GRY)
    rect(s, x, y2, Inches(0.08), ch2, ACCENT)
    text(s, x + Inches(0.18), y2 + Inches(0.12), cw2 - Inches(0.25), Inches(0.4),
         t, size=11, bold=True, color=NAVY)
    text(s, x + Inches(0.18), y2 + Inches(0.5), cw2 - Inches(0.25), ch2 - Inches(0.55),
         b, size=10, color=GRAY_DARK)

# ============================================================
# SLIDE 6 — DRIFT, BASELINES & CONFIG SYNC
# ============================================================
s = add_blank()
page_header(s, "Drift, Baselines & Config Sync",
            "Know exactly what changed, by whom, against which baseline")

# left columns
cols = [
    ("Configuration Baselines",
     ["Capture known-good state per cloud account",
      "Versioned snapshots (BaselineVersion model)",
      "Per-resource fingerprint diffing",
      "Compare any two versions side-by-side"],
     ACCENT),
    ("Config Change Monitor",
     ["Streaming detection of resource changes",
      "Classified: security-impact vs cosmetic",
      "Who-what-when retained per change",
      "Replayable timeline per resource"],
     ORANGE),
    ("Freeze Windows",
     ["Block remediation in release windows",
      "Per-account / per-cloud / global",
      "Audit trail of who declared the freeze",
      "Live banner in dashboard"],
     TEAL),
]
y = Inches(1.15); cw = Inches(4.1); ch = Inches(2.7); gap = Inches(0.1)
for i, (t, items, c) in enumerate(cols):
    x = Inches(0.45) + (cw + gap) * i
    rounded(s, x, y, cw, ch, WHITE, line=LIGHT_GRY)
    rect(s, x, y, cw, Inches(0.5), c)
    text(s, x, y + Inches(0.1), cw, Inches(0.32),
         t, size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    yy = y + Inches(0.65)
    for line in items:
        text(s, x + Inches(0.2), yy, cw - Inches(0.3), Inches(0.4),
             "▸  " + line, size=10, color=GRAY_DARK)
        yy += Inches(0.4)

# bottom: drift mock screenshot
mock_drift(s, Inches(0.45), Inches(4.05), Inches(12.5), Inches(3.05))
text(s, Inches(0.45), SH - Inches(0.35), Inches(12), Inches(0.22),
     "↑ Baseline Drift view — side-by-side compare with severity-tagged deltas",
     size=8, color=GRAY_MID)

# ============================================================
# SLIDE 7 — REPORTING, ALERTS & INTEGRATIONS
# ============================================================
s = add_blank()
page_header(s, "Reporting, Alerts & Integrations",
            "Right finding · right team · right time")

# left: reporting
rounded(s, Inches(0.45), Inches(1.15), Inches(4.1), Inches(2.9), WHITE, line=LIGHT_GRY)
rect(s, Inches(0.45), Inches(1.15), Inches(4.1), Inches(0.5), DEEP_BLUE)
text(s, Inches(0.45), Inches(1.23), Inches(4.1), Inches(0.4),
     "Reporting", size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
rb = ["Per-account / per-cloud / per-framework",
      "Formats: PDF · HTML · CSV · JSON",
      "Scheduled (cron) + on-demand",
      "ReportSchedule + ReportRun audit trail",
      "Posture history exported for board"]
yy = Inches(1.75)
for line in rb:
    text(s, Inches(0.65), yy, Inches(3.85), Inches(0.38),
         "▸  " + line, size=10, color=GRAY_DARK); yy += Inches(0.4)

# middle: alerts
rounded(s, Inches(4.65), Inches(1.15), Inches(4.1), Inches(2.9), WHITE, line=LIGHT_GRY)
rect(s, Inches(4.65), Inches(1.15), Inches(4.1), Inches(0.5), RED)
text(s, Inches(4.65), Inches(1.23), Inches(4.1), Inches(0.4),
     "Alerts", size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
al = ["Slack channel routing by severity",
      "Email digest or per-event",
      "AlertLog — dedup + throttle aware",
      "Freeze-window aware (no noise)",
      "Quiet-hours configurable"]
yy = Inches(1.75)
for line in al:
    text(s, Inches(4.85), yy, Inches(3.85), Inches(0.38),
         "▸  " + line, size=10, color=GRAY_DARK); yy += Inches(0.4)

# right: integrations
rounded(s, Inches(8.85), Inches(1.15), Inches(4.1), Inches(2.9), WHITE, line=LIGHT_GRY)
rect(s, Inches(8.85), Inches(1.15), Inches(4.1), Inches(0.5), TEAL)
text(s, Inches(8.85), Inches(1.23), Inches(4.1), Inches(0.4),
     "Integrations", size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
ig = ["Jira / ServiceNow ticketing",
      "Splunk / DataDog / SIEM webhooks",
      "Slack · Email · generic webhook",
      "Per-finding routing rules",
      "IntegrationLog for audit"]
yy = Inches(1.75)
for line in ig:
    text(s, Inches(9.05), yy, Inches(3.85), Inches(0.38),
         "▸  " + line, size=10, color=GRAY_DARK); yy += Inches(0.4)

# bottom: reports mock screenshot
mock_reports(s, Inches(0.45), Inches(4.2), Inches(12.5), Inches(2.95))
text(s, Inches(0.45), SH - Inches(0.35), Inches(12), Inches(0.22),
     "↑ Findings/Reports view — filters, severity-aware rows, PDF/CSV/JSON export, inline status",
     size=8, color=GRAY_MID)

# ============================================================
# SLIDE 8 — CLOSING / REPORT SUMMARY
# ============================================================
s = add_blank()
rect(s, 0, 0, SW, SH, NAVY)
acc = s.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, Inches(8.3), 0, Inches(5.0), SH)
acc.fill.solid(); acc.fill.fore_color.rgb = DEEP_BLUE; acc.line.fill.background()
rect(s, 0, Inches(1.7), Inches(0.5), Inches(0.08), ACCENT)
text(s, Inches(0.7), Inches(1.05), Inches(11), Inches(0.4),
     "SUMMARY", size=13, bold=True, color=ACCENT)
text(s, Inches(0.7), Inches(1.4), Inches(11), Inches(1.0),
     "One platform.  Three clouds.\nZero per-asset pricing.",
     size=34, bold=True, color=WHITE)

# 4 mini cards
caps = [
    ("70+",  "Scanners"),
    ("6",    "Frameworks"),
    ("3",    "Clouds"),
    ("31",   "UI Pages"),
]
y = Inches(3.4); w = Inches(1.8); h = Inches(1.2); gap = Inches(0.15)
for i, (n, l) in enumerate(caps):
    x = Inches(0.7) + (w + gap) * i
    rounded(s, x, y, w, h, ACCENT)
    text(s, x, y + Inches(0.12), w, Inches(0.6),
         n, size=28, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    text(s, x, y + Inches(0.7), w, Inches(0.4),
         l, size=10, color=WHITE, align=PP_ALIGN.CENTER)

# value bullets
yv = Inches(4.95)
vals = [
    "▸  Replaces Wiz / Orca / Prisma + GuardDuty + Inspector + Macie",
    "▸  Detect → Approve → Auto-remediate, end to end",
    "▸  Self-hosted — your data, your VPC, predictable cost",
    "▸  Up and running in an afternoon (docker-compose up)",
]
for line in vals:
    text(s, Inches(0.7), yv, Inches(11.5), Inches(0.32),
         line, size=13, color=WHITE); yv += Inches(0.35)

# CTAs
rounded(s, Inches(0.7), SH - Inches(0.85), Inches(2.6), Inches(0.5), ACCENT)
text(s, Inches(0.7), SH - Inches(0.77), Inches(2.6), Inches(0.4),
     "REQUEST DEMO", size=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
rounded(s, Inches(3.45), SH - Inches(0.85), Inches(2.6), Inches(0.5), WHITE)
text(s, Inches(3.45), SH - Inches(0.77), Inches(2.6), Inches(0.4),
     "VIEW DOCS", size=11, bold=True, color=NAVY, align=PP_ALIGN.CENTER)

# ---------- save ----------
out = "/mnt/d/Data/Linux-Data/IML/vul-scanner/aws-scanner/docs/CSPM_Platform_Showcase_Short.pptx"
prs.save(out)
print("Saved:", out)
print("Slides:", len(prs.slides.__iter__.__self__._sldIdLst))
